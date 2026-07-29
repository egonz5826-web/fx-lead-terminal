#!/usr/bin/env python3
"""One-time import of Jesus's color-coded Excel client book into Supabase.
Reads cell fill colors (theme-resolved) to map each client to a status.
Usage: python3 scripts/import_client_book.py "/path/to/Clientes Chucho.xlsx"
Re-runnable: clears owner=jesus rows first."""
import sys, re, json, colorsys, urllib.request
import openpyxl

XLSX = sys.argv[1] if len(sys.argv) > 1 else "/Users/emilio/Downloads/Clientes Chucho - Jose recent.xlsx"

env = {}
for line in open(".env.local"):
    m = re.match(r'^([A-Z_]+)=(.*)$', line)
    if m:
        env[m.group(1)] = m.group(2).strip().strip('"').strip("'")
U = env["NEXT_PUBLIC_SUPABASE_URL"]
KEY = env["SUPABASE_SERVICE_ROLE_KEY"]

wb = openpyxl.load_workbook(XLSX)
ws = wb["List of accounts"]

theme_xml = wb.loaded_theme.decode() if isinstance(wb.loaded_theme, bytes) else wb.loaded_theme
block = re.search(r'<a:clrScheme.*?</a:clrScheme>', theme_xml, re.S).group()
pal = [m.group(2) for m in re.finditer(r'<a:(srgbClr|sysClr)[^>]*?(?:val|lastClr)="([0-9A-Fa-f]{6})"', block)][:12]

def theme_rgb(idx):
    swap = {0: 1, 1: 0, 2: 3, 3: 2}
    i = swap.get(idx, idx)
    return pal[i] if i < len(pal) else "FFFFFF"

def apply_tint(hexc, tint):
    def t(v):
        v /= 255
        v = v * (1 + tint) if tint < 0 else v * (1 - tint) + tint
        return max(0, min(255, round(v * 255)))
    return (t(int(hexc[0:2], 16)), t(int(hexc[2:4], 16)), t(int(hexc[4:6], 16)))

def rgb_of(cell):
    f = cell.fill
    if not f or f.patternType is None:
        return None
    c = f.fgColor
    if c.type == "rgb" and c.rgb and c.rgb != "00000000":
        h = c.rgb[-6:]
        return (int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))
    if c.type == "theme":
        return apply_tint(theme_rgb(c.theme), float(c.tint or 0))
    return None

def classify(rgb):
    if rgb is None:
        return "Uncategorized"
    r, g, b = rgb
    h, s, v = colorsys.rgb_to_hsv(r / 255, g / 255, b / 255)
    H = h * 360
    if s < 0.12 and v > 0.85:
        return "Open and trading"
    if s < 0.12:
        return "Uncategorized"
    if H < 15 or H >= 345:
        return "VIP / priority" if v >= 0.85 else "Cancelled"
    if 15 <= H < 45:
        return "On hold by Management (Mexican corp)"
    if 45 <= H < 70:
        return "On hold by client"
    if 70 <= H < 160:
        return "Open and trading"
    if 160 <= H < 200:
        return "On SF"
    if 200 <= H < 255:
        return "Chucho needs to follow up"
    if 255 <= H < 300:
        return "Open but not trading"
    return "Uncategorized"

def cell(r, c):
    v = ws.cell(r, c).value
    return str(v).strip() if v is not None and str(v).strip() else None

rows = []
for r in range(2, ws.max_row + 1):
    name = cell(r, 5)
    if not name:
        continue
    jh, jr = cell(r, 16), cell(r, 15)
    comments = " | ".join(p for p in [f"JH: {jh}" if jh else None, f"JR: {jr}" if jr else None] if p)
    num = ws.cell(r, 4).value
    rows.append({
        "number": int(num) if isinstance(num, (int, float)) else None,
        "name": name,
        "status": classify(rgb_of(ws.cell(r, 5))),
        "on_sf": (cell(r, 6) or "").lower() == "x",
        "consultant": cell(r, 7),
        "website": cell(r, 8),
        "referenciador": cell(r, 9),
        "communication": cell(r, 10),
        "dealer": cell(r, 11),
        "traded": cell(r, 12),
        "last_follow_up": cell(r, 13),
        "follow_up": cell(r, 14),
        "comments": comments or None,
        "contact_rr": cell(r, 17),
        "owner": "jesus",
    })

def req(method, path, body=None):
    r = urllib.request.Request(
        f"{U}/rest/v1/{path}", method=method,
        data=json.dumps(body).encode() if body is not None else None,
        headers={"apikey": KEY, "Authorization": f"Bearer {KEY}",
                 "Content-Type": "application/json", "Prefer": "return=minimal"})
    return urllib.request.urlopen(r)

# clear existing jesus rows, then insert fresh
try:
    req("DELETE", "client_book?owner=eq.jesus")
except Exception as e:
    print("clear warning:", e)
for i in range(0, len(rows), 100):
    req("POST", "client_book", rows[i:i + 100])
print(f"Imported {len(rows)} clients for jesus.")
from collections import Counter
for s, c in Counter(x["status"] for x in rows).most_common():
    print(f"  {c:3d}  {s}")
